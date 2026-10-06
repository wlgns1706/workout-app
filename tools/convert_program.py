"""엑셀 프로그램 파일을 앱이 읽는 JSON으로 바꾼다.

사용법:
  python tools/convert_program.py <엑셀 경로> <출력 경로> --id base-one --name "근비대 One"
"""
import argparse
import json
import re
import zipfile
import sys
import warnings

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

warnings.filterwarnings("ignore")
LABEL = re.compile(r"^W(\d)D(\d)")


def clean(value):
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        return value.strip()
    return value


def extension_lists(path):
    """엑셀 확장 형식(x14)으로 저장된 드롭다운. openpyxl이 읽지 못해서 XML에서 직접 읽는다.
    시트 이름 → [(목록 범위 수식, 적용 칸 범위)]"""
    z = zipfile.ZipFile(path)
    workbook = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    targets = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
    targets.update({k: v for v, k in re.findall(r'Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
    result = {}
    for name, rid in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', workbook):
        target = targets[rid].lstrip("/")
        xml = z.read(target if target.startswith("xl/") else "xl/" + target).decode("utf-8")
        result[name] = re.findall(
            r'<x14:dataValidation [^>]*type="list".*?<xm:f>(.*?)</xm:f>.*?<xm:sqref>(.*?)</xm:sqref>', xml, re.S
        )
    return result


def alternatives_of(ws, wb, extra):
    """종목 칸의 드롭다운 목록. (행, 열) → 대체 운동 이름 목록"""
    result = {}
    lists = [(dv.formula1, str(dv.sqref)) for dv in ws.data_validations.dataValidation if dv.type == "list" and dv.formula1]
    lists += [(f.replace("&apos;", "'"), sq) for f, sq in extra]
    for formula, sqref in lists:
        if "!" not in formula:
            continue
        sheet_name, ref = formula.rsplit("!", 1)
        source = wb[sheet_name.strip("'")]
        min_col, min_row, max_col, max_row = range_boundaries(ref.replace("$", ""))
        names = [
            str(source.cell(r, c).value).strip()
            for r in range(min_row, max_row + 1)
            for c in range(min_col, max_col + 1)
            if source.cell(r, c).value not in (None, "")
        ]
        for cell_range in sqref.split():
            c1, r1, c2, r2 = range_boundaries(cell_range)
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    result[(r, c)] = names
    return result


def parse_sheet(ws, wb, extra):
    """시트 하나(블록)를 4주 목록으로 바꾼다."""
    alts = alternatives_of(ws, wb, extra)
    labels = {}
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and LABEL.match(cell.value.strip()):
                labels.setdefault(cell.column, []).append((cell.row, cell.value.strip()))

    # 주차는 글자(W1~W4)가 아니라 열의 순서로 정한다. 원본 엑셀에 주차 글자가 잘못 적힌 칸이 있다.
    week_of_col = {col: i + 1 for i, col in enumerate(sorted(labels))}
    weeks = {}
    for col, items in labels.items():
        items.sort()
        for i, (top, text) in enumerate(items):
            bottom = items[i + 1][0] if i + 1 < len(items) else ws.max_row + 1
            match = LABEL.match(text)
            day = {
                "dayNo": int(match.group(2)),
                "optional": "선택" in text,
                "exercises": [],
                "cardio": None,
                "optionsText": None,
            }
            current = None
            for r in range(top, bottom):
                name = clean(ws.cell(r, col + 1).value)
                sets = clean(ws.cell(r, col + 2).value)
                reps = clean(ws.cell(r, col + 3).value)
                rpe = clean(ws.cell(r, col + 4).value)
                if isinstance(name, str) and name.startswith("옵션"):
                    day["optionsText"] = name
                    current = None
                    continue
                if isinstance(name, str) and "유산소" in name:
                    day["cardio"] = {
                        "label": name,
                        "detail": "" if sets is None else str(sets),
                        "required": "필수" in name,
                    }
                    current = None
                    continue
                if name == "Notes":
                    current = None
                    continue
                if name:
                    current = {"name": str(name), "rows": []}
                    options = [a for a in alts.get((r, col + 1), []) if a != str(name)]
                    if options:
                        current["alternatives"] = options
                    day["exercises"].append(current)
                if current is not None and isinstance(sets, int):
                    current["rows"].append(
                        {"sets": sets, "reps": "" if reps is None else str(reps), "rpe": rpe or None}
                    )
            day["exercises"] = [e for e in day["exercises"] if e["rows"]]
            weeks.setdefault(week_of_col[col], []).append(day)

    result = []
    for week_no in sorted(weeks):
        days = sorted(weeks[week_no], key=lambda d: d["dayNo"])
        rest = all(not d["exercises"] and not d["optionsText"] for d in days)
        result.append({"rest": rest, "days": days})
    return result


def read_chart(wb):
    ws = wb["RPE Chart"]
    reps = [int(ws.cell(2, c).value) for c in range(4, 16)]
    rows = [
        {"rpe": float(ws.cell(r, 3).value), "pct": [float(ws.cell(r, c).value) for c in range(4, 16)]}
        for r in range(3, 11)
    ]
    return {"reps": reps, "rows": rows}


def convert(path, program_id, name):
    wb = load_workbook(path, data_only=True)
    extra = extension_lists(path)
    blocks = [{"weeks": parse_sheet(ws, wb, extra.get(ws.title, []))} for ws in wb.worksheets if ws.title.startswith("프로그램")]
    if not blocks:
        sys.exit("'프로그램'으로 시작하는 시트가 없습니다.")
    for b, block in enumerate(blocks, 1):
        if len(block["weeks"]) != 4:
            sys.exit(f"블록 {b}: 4주가 아닙니다 ({len(block['weeks'])}주).")
        for w, week in enumerate(block["weeks"], 1):
            if len(week["days"]) != 6:
                sys.exit(f"블록 {b} {w}주차: 요일이 6개가 아닙니다 ({len(week['days'])}개).")
    return {
        "type": "workout-program",
        "formatVersion": 1,
        "id": program_id,
        "name": name,
        "blocks": blocks,
        "rpeChart": read_chart(wb),
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx")
    parser.add_argument("out")
    parser.add_argument("--id", required=True)
    parser.add_argument("--name", required=True)
    args = parser.parse_args()
    program = convert(args.xlsx, args.id, args.name)
    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(program, f, ensure_ascii=False)
    print(f"저장: {args.out}")

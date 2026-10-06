"""변환한 프로그램 파일이 엑셀과 맞는지 확인한다.

사용법:
  python tools/check_convert.py <엑셀 경로> <프로그램 JSON 경로>
"""
import json
import re
import sys
import warnings

from openpyxl import load_workbook

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
LABEL = re.compile(r"^W(\d)D(\d)")


def count_in_sheet(ws):
    """요일 글자 개수와, Sets 열에 숫자가 있는 줄 수를 센다."""
    label_cols = set()
    labels = 0
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and LABEL.match(cell.value.strip()):
                labels += 1
                label_cols.add(cell.column)
    set_rows = 0
    for col in label_cols:
        for r in range(1, ws.max_row + 1):
            if isinstance(ws.cell(r, col + 2).value, (int, float)):
                set_rows += 1
    return labels, set_rows


def main(xlsx, out):
    wb = load_workbook(xlsx, data_only=True)
    program = json.load(open(out, encoding="utf-8"))
    sheets = [ws for ws in wb.worksheets if ws.title.startswith("프로그램")]
    assert len(program["blocks"]) == len(sheets) == 3, "블록은 3개여야 한다"
    for ws, block in zip(sheets, program["blocks"]):
        labels, set_rows = count_in_sheet(ws)
        days = [d for w in block["weeks"] for d in w["days"]]
        rows = sum(len(e["rows"]) for d in days for e in d["exercises"])
        assert labels == len(days) == 24, f"{ws.title}: 요일 수 {labels} / {len(days)}"
        assert set_rows == rows, f"{ws.title}: 세트 줄 수가 다르다 (엑셀 {set_rows}, 변환 {rows})"
        for w in block["weeks"]:
            assert [d["dayNo"] for d in w["days"]] == [1, 2, 3, 4, 5, 6]
            assert [d["optional"] for d in w["days"]] == [False, False, False, False, True, True]
        print(f"{ws.title}: 요일 {len(days)}, 세트 줄 {rows} 일치")
    assert program["blocks"][2]["weeks"][3]["rest"] is True, "블록 3의 4주차는 완전 휴식이어야 한다"
    assert not any(w["rest"] for b in program["blocks"][:2] for w in b["weeks"]), "블록 1, 2에는 휴식 주가 없어야 한다"
    week1 = program["blocks"][0]["weeks"][0]
    assert all(d["cardio"] for d in week1["days"][:4]), "D1~D4에는 유산소 줄이 있어야 한다"
    assert all(d["optionsText"] for d in week1["days"][4:]), "D5, D6에는 옵션 글이 있어야 한다"
    with_alts = [e for b in program["blocks"] for w in b["weeks"] for d in w["days"] for e in d["exercises"] if e.get("alternatives")]
    assert with_alts, "대체 운동 목록이 하나도 없다"
    assert all(e["name"] not in e["alternatives"] for e in with_alts), "대체 운동 목록에 자기 자신이 들어 있다"
    print(f"대체 운동이 있는 종목 {len(with_alts)}개, 예: {with_alts[0]['name']} → {with_alts[0]['alternatives']}")
    assert program["rpeChart"]["reps"] == list(range(1, 13))
    assert len(program["rpeChart"]["rows"]) == 8
    assert program["rpeChart"]["rows"][0]["pct"][0] == 1
    print("확인 완료")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
